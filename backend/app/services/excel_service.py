import io
import os
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.services.onedrive_service import (
    obtener_excel_onedrive,
    estado_fuente_onedrive,
    subir_excel_onedrive,
    invalidar_cache_onedrive,
    actualizar_registro
)

BASE_DIR = Path(__file__).resolve().parent.parent
RUTA_EXCEL_LOCAL = BASE_DIR / "data" / "INVENTARIO GENERAL TI - ACTUAL.xlsx"

_KPI_CACHE = {
    "etag": None,
    "data": None,
}

_HOJAS_CACHE = {
    "etag": None,
    "hojas": {},
}


def _data_source() -> str:
    return os.getenv("DATA_SOURCE", "local").strip().lower()

def invalidar_caches_excel():
    global _KPI_CACHE
    global _HOJAS_CACHE

    _KPI_CACHE = {
        "etag": None,
        "data": None,
    }

    _HOJAS_CACHE = {
        "etag": None,
        "hojas": {},
    }

    if _data_source() == "onedrive":
        invalidar_cache_onedrive()

def obtener_fuente_excel():
    """Retorna una fuente compatible con pandas.read_excel.

    DATA_SOURCE=local     -> usa el Excel incluido en backend/app/data
    DATA_SOURCE=onedrive  -> descarga/usa caché del Excel privado en Microsoft 365
    """
    source = _data_source()

    if source == "onedrive":
        return obtener_excel_onedrive()

    if source == "local":
        if not RUTA_EXCEL_LOCAL.exists():
            raise FileNotFoundError(
                f"No se encontró el archivo Excel local en: {RUTA_EXCEL_LOCAL}"
            )
        return RUTA_EXCEL_LOCAL

    raise RuntimeError(
        f"DATA_SOURCE='{source}' no es válido. Usa 'local' u 'onedrive'."
    )

def _obtener_dataframe_hoja(nombre_hoja):
    global _HOJAS_CACHE
    global _KPI_CACHE

    source = _data_source()

    if source == "onedrive":
        estado = estado_fuente_onedrive()
        etag_actual = estado.get("etag")

        # Si cambió el Excel, invalidamos todas las cachés.
        if _HOJAS_CACHE["etag"] != etag_actual:
            _HOJAS_CACHE = {
                "etag": etag_actual,
                "hojas": {},
            }

            _KPI_CACHE = {
                "etag": None,
                "data": None,
            }

        # Si la hoja ya fue procesada para esta versión,
        # la devolvemos directamente.
        if nombre_hoja in _HOJAS_CACHE["hojas"]:
            return _HOJAS_CACHE["hojas"][nombre_hoja]

        # Si no está cacheada, la leemos una sola vez.
        df = pd.read_excel(
            obtener_fuente_excel(),
            sheet_name=nombre_hoja,
        )

        df = df.fillna("")

        _HOJAS_CACHE["hojas"][nombre_hoja] = df

        return df

    # Modo local
    df = pd.read_excel(
        obtener_fuente_excel(),
        sheet_name=nombre_hoja,
    )

    return df.fillna("")

def obtener_hoja(nombre_hoja):
    try:
        df = _obtener_dataframe_hoja(nombre_hoja)

        registros = df.to_dict(
            orient="records"
        )

        # Pandas usa índice 0 para la primera fila de datos.
        # En Excel:
        # fila 1 = encabezados
        # fila 2 = primer registro
        #
        # Guardamos este dato únicamente como identificador técnico.
        for indice, registro in enumerate(
            registros,
            start=2,
        ):
            registro["__row_id"] = indice

        return registros

    except ValueError as e:
        return {
            "error":
                f"No se encontró la hoja "
                f"'{nombre_hoja}': {str(e)}"
        }

    except Exception as e:
        return {
            "error": str(e)
        }


def obtener_kpis_principales():
    global _KPI_CACHE

    hojas = {
        "laptops": "LAPTOPS",
        "celulares": "CELULARES",
        "monitores": "MONITORES",
        "impresoras": "IMPRESORAS",
        "chips": "ASIGNACIÓN CHIPS",
        "modem": "MODEM",
        "exchange": "EXCHANGE",
        "reportados": "EQUIPOS REPORTADOS",
    }

    resultado = {}
    total = 0

    try:
        if _data_source() == "onedrive":
            estado = estado_fuente_onedrive()
            etag_actual = estado.get("etag")

            if (
                etag_actual
                and _KPI_CACHE["etag"] == etag_actual
                and _KPI_CACHE["data"] is not None
            ):
                return _KPI_CACHE["data"]

        else:
            etag_actual = None

        for clave, hoja in hojas.items():
            try:
                df = _obtener_dataframe_hoja(hoja)

                cantidad = len(df)

                resultado[clave] = cantidad
                total += cantidad

            except Exception as e:
                resultado[clave] = 0
                resultado[f"{clave}_error"] = str(e)

        resultado["total"] = total

        if _data_source() == "onedrive":
            _KPI_CACHE = {
                "etag": etag_actual,
                "data": resultado.copy(),
            }

        return resultado

    except Exception as e:
        resultado["error_fuente"] = str(e)

        for clave in hojas:
            resultado[clave] = 0

        resultado["total"] = 0

        return resultado
    

def obtener_estado_fuente():
    source = _data_source()

    if source == "onedrive":
        try:
            return estado_fuente_onedrive()
        except Exception as e:
            return {"fuente": "onedrive", "estado": "error", "error": str(e)}

    return {
        "fuente": "local",
        "archivo": RUTA_EXCEL_LOCAL.name,
        "existe": RUTA_EXCEL_LOCAL.exists(),
    }

def obtener_etag_actual():
    if _data_source() != "onedrive":
        return None

    try:
        estado = estado_fuente_onedrive()
        return estado.get("etag")
    except Exception:
        return None

def actualizar_registro(
    nombre_hoja: str,
    fila_excel: int,
    nuevos_datos: dict,
    etag_esperado: str | None = None,
):
    """
    Actualiza una fila existente del Excel.

    nombre_hoja:
        Nombre exacto de la hoja.

    fila_excel:
        Número real de fila dentro del Excel.

    nuevos_datos:
        Diccionario:
        {
            "USUARIO": "...",
            "CARGO": "...",
            ...
        }

    etag_esperado:
        eTag que tenía el archivo antes de editar.
        Sirve para evitar sobrescribir cambios externos.
    """

    if not nombre_hoja:
        raise ValueError(
            "Debe especificarse la hoja."
        )

    if not isinstance(fila_excel, int):
        raise ValueError(
            "fila_excel debe ser un número entero."
        )

    if fila_excel < 2:
        raise ValueError(
            "No se puede modificar la fila de encabezados."
        )

    if not isinstance(nuevos_datos, dict):
        raise ValueError(
            "Los datos deben enviarse como un objeto."
        )

    if not nuevos_datos:
        raise ValueError(
            "No se recibieron campos para actualizar."
        )

    source = _data_source()

    if source != "onedrive":
        raise RuntimeError(
            "La edición actualmente está habilitada "
            "solo para DATA_SOURCE=onedrive."
        )

    # --------------------------------------------------
    # 1. Descargar siempre la versión actual
    # --------------------------------------------------

    excel_bytes = obtener_excel_onedrive(
        force_refresh=True
    )

    # --------------------------------------------------
    # 2. Abrir workbook
    # --------------------------------------------------

    workbook = load_workbook(
        excel_bytes
    )

    if nombre_hoja not in workbook.sheetnames:
        raise ValueError(
            f"No existe la hoja '{nombre_hoja}'."
        )

    worksheet = workbook[nombre_hoja]

    # --------------------------------------------------
    # 3. Leer encabezados de Excel
    # --------------------------------------------------

    encabezados = {}

    for columna in range(
        1,
        worksheet.max_column + 1
    ):

        valor = worksheet.cell(
            row=1,
            column=columna
        ).value

        if valor is None:
            continue

        encabezados[
            str(valor).strip()
        ] = columna

    # --------------------------------------------------
    # 4. Validar fila
    # --------------------------------------------------

    if fila_excel > worksheet.max_row:
        raise ValueError(
            f"La fila {fila_excel} ya no existe "
            f"en la hoja '{nombre_hoja}'."
        )

    # --------------------------------------------------
    # 5. Modificar únicamente las columnas recibidas
    # --------------------------------------------------

    campos_actualizados = []

    for campo, nuevo_valor in nuevos_datos.items():

        # Campos internos del sistema
        if campo.startswith("__"):
            continue

        if campo not in encabezados:
            continue

        numero_columna = encabezados[campo]

        worksheet.cell(
            row=fila_excel,
            column=numero_columna
        ).value = nuevo_valor

        campos_actualizados.append(
            campo
        )

    if not campos_actualizados:
        raise ValueError(
            "Ninguno de los campos enviados "
            "existe en la hoja de Excel."
        )

    # --------------------------------------------------
    # 6. Guardar workbook en memoria
    # --------------------------------------------------

    salida = io.BytesIO()

    workbook.save(
        salida
    )

    contenido_actualizado = (
        salida.getvalue()
    )

    # --------------------------------------------------
    # 7. Subir Excel a SharePoint
    # --------------------------------------------------

    resultado = subir_excel_onedrive(
        contenido=contenido_actualizado,
        etag_esperado=etag_esperado,
    )

    # --------------------------------------------------
    # 8. Invalidar cache
    # --------------------------------------------------

    invalidar_caches_excel()

    return {
        "ok": True,
        "mensaje":
            "Registro actualizado correctamente.",
        "hoja": nombre_hoja,
        "fila_excel": fila_excel,
        "campos_actualizados":
            campos_actualizados,
        "archivo":
            resultado.get("archivo"),
        "etag":
            resultado.get("etag"),
    }

def actualizar_registro(
    nombre_hoja: str,
    fila_excel: int,
    nuevos_datos: dict,
    etag_esperado: str | None = None,
):
    import io

    from openpyxl import load_workbook

    from app.services.onedrive_service import (
        obtener_excel_onedrive,
        subir_excel_onedrive,
        invalidar_cache_onedrive,
    )

    global _KPI_CACHE
    global _HOJAS_CACHE

    if not nombre_hoja:
        raise ValueError(
            "Debe especificarse la hoja."
        )

    if not isinstance(fila_excel, int):
        raise ValueError(
            "fila_excel debe ser un número entero."
        )

    if fila_excel < 2:
        raise ValueError(
            "No se puede modificar la fila de encabezados."
        )

    if not isinstance(nuevos_datos, dict):
        raise ValueError(
            "Los datos deben enviarse como un objeto."
        )

    if not nuevos_datos:
        raise ValueError(
            "No se recibieron campos para actualizar."
        )

    if _data_source() != "onedrive":
        raise RuntimeError(
            "La edición está habilitada solo "
            "para DATA_SOURCE=onedrive."
        )

    # Obtener siempre la versión más reciente
    excel_bytes = obtener_excel_onedrive(
        force_refresh=True
    )

    # Abrir el Excel manteniendo sus hojas y formatos
    workbook = load_workbook(
        excel_bytes
    )

    if nombre_hoja not in workbook.sheetnames:
        raise ValueError(
            f"No existe la hoja '{nombre_hoja}'."
        )

    worksheet = workbook[nombre_hoja]

    # Obtener encabezados de la fila 1
    encabezados = {}

    for columna in range(
        1,
        worksheet.max_column + 1
    ):
        valor = worksheet.cell(
            row=1,
            column=columna
        ).value

        if valor is None:
            continue

        encabezados[
            str(valor).strip()
        ] = columna

    if fila_excel > worksheet.max_row:
        raise ValueError(
            f"La fila {fila_excel} no existe "
            f"en la hoja '{nombre_hoja}'."
        )

    campos_actualizados = []

    # Actualizar solamente los campos recibidos
    for campo, nuevo_valor in nuevos_datos.items():

        # Ignorar identificadores internos
        if campo.startswith("__"):
            continue

        if campo not in encabezados:
            continue

        numero_columna = encabezados[campo]

        worksheet.cell(
            row=fila_excel,
            column=numero_columna
        ).value = nuevo_valor

        campos_actualizados.append(
            campo
        )

    if not campos_actualizados:
        raise ValueError(
            "Ninguno de los campos enviados "
            "existe en la hoja."
        )

    # Guardar Excel en memoria
    salida = io.BytesIO()

    workbook.save(
        salida
    )

    contenido_actualizado = salida.getvalue()

    # Subir Excel actualizado a SharePoint
    resultado = subir_excel_onedrive(
        contenido=contenido_actualizado,
        etag_esperado=etag_esperado,
    )

    # Limpiar cachés
    _KPI_CACHE = {
        "etag": None,
        "data": None,
    }

    _HOJAS_CACHE = {
        "etag": None,
        "hojas": {},
    }

    invalidar_cache_onedrive()

    return {
        "ok": True,
        "mensaje": "Registro actualizado correctamente.",
        "hoja": nombre_hoja,
        "fila_excel": fila_excel,
        "campos_actualizados": campos_actualizados,
        "archivo": resultado.get("archivo"),
        "etag": resultado.get("etag"),
    }


def actualizar_registro(
    nombre_hoja: str,
    fila_excel: int,
    nuevos_datos: dict,
    etag_esperado: str | None = None,
):
    import io
    from openpyxl import load_workbook

    from app.services.onedrive_service import (
        obtener_excel_onedrive,
        subir_excel_onedrive,
        invalidar_cache_onedrive,
    )

    global _KPI_CACHE
    global _HOJAS_CACHE

    if not nombre_hoja:
        raise ValueError(
            "Debe especificarse la hoja."
        )

    if not isinstance(fila_excel, int):
        raise ValueError(
            "fila_excel debe ser un número entero."
        )

    if fila_excel < 2:
        raise ValueError(
            "No se puede modificar la fila de encabezados."
        )

    if not isinstance(nuevos_datos, dict):
        raise ValueError(
            "Los datos deben enviarse como un objeto."
        )

    if not nuevos_datos:
        raise ValueError(
            "No se recibieron campos para actualizar."
        )

    if _data_source() != "onedrive":
        raise RuntimeError(
            "La edición está habilitada solo para DATA_SOURCE=onedrive."
        )

    # Descargar siempre la versión más reciente
    excel_bytes = obtener_excel_onedrive(
        force_refresh=True
    )

    workbook = load_workbook(
        excel_bytes
    )

    if nombre_hoja not in workbook.sheetnames:
        raise ValueError(
            f"No existe la hoja '{nombre_hoja}'."
        )

    worksheet = workbook[nombre_hoja]

    # Obtener encabezados de la fila 1
    encabezados = {}

    for columna in range(
        1,
        worksheet.max_column + 1
    ):
        valor = worksheet.cell(
            row=1,
            column=columna
        ).value

        if valor is None:
            continue

        encabezados[
            str(valor).strip()
        ] = columna

    if fila_excel > worksheet.max_row:
        raise ValueError(
            f"La fila {fila_excel} no existe en la hoja '{nombre_hoja}'."
        )

    campos_actualizados = []

    for campo, nuevo_valor in nuevos_datos.items():

        if campo.startswith("__"):
            continue

        if campo not in encabezados:
            continue

        numero_columna = encabezados[campo]

        worksheet.cell(
            row=fila_excel,
            column=numero_columna
        ).value = nuevo_valor

        campos_actualizados.append(
            campo
        )

    if not campos_actualizados:
        raise ValueError(
            "Ninguno de los campos enviados existe en la hoja."
        )

    salida = io.BytesIO()

    workbook.save(
        salida
    )

    contenido_actualizado = salida.getvalue()

    resultado = subir_excel_onedrive(
        contenido=contenido_actualizado,
        etag_esperado=etag_esperado,
    )

    # Limpiar cachés
    _KPI_CACHE = {
        "etag": None,
        "data": None,
    }

    _HOJAS_CACHE = {
        "etag": None,
        "hojas": {},
    }

    invalidar_cache_onedrive()

    return {
        "ok": True,
        "mensaje": "Registro actualizado correctamente.",
        "hoja": nombre_hoja,
        "fila_excel": fila_excel,
        "campos_actualizados": campos_actualizados,
        "archivo": resultado.get("archivo"),
        "etag": resultado.get("etag"),
    }